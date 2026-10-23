"""Render calibration_parameters.yaml as a Word document (landscape A4, one table per section).

Called by render_table.py, which writes the HTML and the .docx together; can also run alone:

    uv run python european_values/render_docx.py

The same cells as the HTML table: parameter (symbol, id, code location), meaning, US value and
source, EU value with its treatment, EU source, justification, status with any to-do. Key
differences carry a star, discuss rows are tinted amber, and LaTeX in the YAML ($\\bar r$,
$\\times$, ...) is turned into plain Unicode since Word has no KaTeX.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from render_table import COLUMNS, SRC, STATUSES, by_section, clean, parse_blocks, rows_for  # noqa: E402

# Column widths in cm; they sum to the usable width of a landscape A4 page with 1.2 cm margins.
WIDTHS = {
    "Parameter": 3.0,
    "Meaning": 3.2,
    "US value": 1.8,
    "US source": 3.0,
    "EU value": 2.0,
    "EU source / hint": 5.3,
    "Justification": 7.0,
    "Status": 2.0,
}
HEADER_FILL = "E8ECF1"
STATUS_FILL = {"decided": "E4F3EA", "discuss": "FDF1DC", "open": "FBE7E5"}
STATUS_INK = {"decided": "2E8B57", "discuss": "B8701A", "open": "B9433E"}
TREATMENT_INK = {"retain": "2E8B57", "recalculate": "B8701A", "recalibrate": "B9433E", "derived": "5F6670"}
DISCUSS_ROW_FILL = "FFF7E8"
MUTED = RGBColor(0x5F, 0x66, 0x70)
LINK = RGBColor(0x2F, 0x5F, 0xA8)

# ------------------------------------------------------------------------------------------------
# LaTeX to plain text
# ------------------------------------------------------------------------------------------------

SYMBOLS = {
    r"\mu": "μ", r"\lambda": "λ", r"\phi": "ϕ", r"\iota": "ι", r"\delta": "δ", r"\sigma": "σ",
    r"\rho": "ρ", r"\psi": "ψ", r"\xi": "ξ", r"\theta": "θ", r"\epsilon": "ε", r"\varepsilon": "ε",
    r"\pi": "π", r"\alpha": "α", r"\beta": "β", r"\gamma": "γ",
    r"\approx": "≈", r"\times": "×", r"\le": "≤", r"\ge": "≥", r"\leq": "≤", r"\geq": "≥",
    r"\cdot": "·", r"\infty": "∞", r"\%": "%", r"\,": " ", r"\;": " ", r"\ ": " ",
}
_FRAC = re.compile(r"\\frac\{([^{}]*)\}\{([^{}]*)\}")
_TEXT = re.compile(r"\\(?:text|mathrm|operatorname)\{([^{}]*)\}")
_BAR = re.compile(r"\\bar\s*(?:\{([^{}]*)\}|(\S))")
_SUB = re.compile(r"_\{([^{}]*)\}")
_SUP = re.compile(r"\^\{([^{}]*)\}")
_MATH = re.compile(r"\$\$?([^$]+)\$\$?")


def _frac(m: re.Match) -> str:
    """a/b, with parentheses around a part that itself contains an operator or a space."""
    num, den = m.group(1).strip(), m.group(2).strip()
    wrap = lambda t: f"({t})" if re.search(r"[ /+\-*×]", t) else t  # noqa: E731
    return f"{wrap(num)}/{wrap(den)}"


def _tex_to_text(tex: str) -> str:
    s = tex
    for _ in range(3):  # nested fractions, a few levels deep
        s = _FRAC.sub(_frac, s)
    s = _TEXT.sub(r"\1", s)
    for k in sorted(SYMBOLS, key=len, reverse=True):  # Greek first, so \bar \mu bars the letter
        s = s.replace(k, SYMBOLS[k])
    s = _BAR.sub(lambda m: (m.group(1) or m.group(2)) + "\u0304", s)
    s = _SUB.sub(r"_\1", s)
    s = _SUP.sub(r"^\1", s)
    s = s.replace("{", "").replace("}", "")
    return " ".join(s.split())


def plain(text: str) -> str:
    """Replace every $...$ in the text by its plain-text rendering."""
    return _MATH.sub(lambda m: _tex_to_text(m.group(1)), text)


# ------------------------------------------------------------------------------------------------
# Word helpers
# ------------------------------------------------------------------------------------------------


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2F5FA8")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(u)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _runs(paragraph, text: str, *, size: float | None = None, italic: bool = False, color: RGBColor | None = None) -> None:
    """Add text to a paragraph, turning bare URLs into hyperlinks."""
    for i, tok in enumerate(text.split(" ")):
        if i:
            paragraph.add_run(" ")
        bare = tok.rstrip(".,;)")
        tail = tok[len(bare):]
        if bare.startswith(("http://", "https://")):
            _hyperlink(paragraph, bare, bare)
            if tail:
                paragraph.add_run(tail)
            continue
        run = paragraph.add_run(tok)
        if size:
            run.font.size = Pt(size)
        run.italic = italic
        if color is not None:
            run.font.color.rgb = color


def _tight(paragraph, *, before: float = 0, after: float = 1, left: float | None = None, hanging: float | None = None) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if left is not None:
        pf.left_indent = Cm(left)
    if hanging is not None:
        pf.first_line_indent = Cm(-hanging)


def _fill_cell(cell, text: str) -> None:
    """Write structured YAML text into a cell: paragraphs, bullets, enumerated items."""
    text = plain(text)
    first = True
    for kind, t in parse_blocks(text):
        if kind == "br":
            continue
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        if kind == "li":
            _runs(p, "•  " + t)
            _tight(p, left=0.35, hanging=0.35)
        elif kind == "en":
            _runs(p, t)
            _tight(p, left=0.35, hanging=0.35)
        else:
            _runs(p, t)
            _tight(p, after=2)
    if first:  # empty text: keep the one empty paragraph tight
        _tight(cell.paragraphs[0])


def _set_width(cell, cm: float) -> None:
    cell.width = Cm(cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    w = OxmlElement("w:tcW")
    w.set(qn("w:w"), str(int(cm * 567)))
    w.set(qn("w:type"), "dxa")
    tc_pr.append(w)


# ------------------------------------------------------------------------------------------------
# Document
# ------------------------------------------------------------------------------------------------


def _page_setup(doc) -> None:
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.2))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(8)
    normal.paragraph_format.space_after = Pt(0)


def _header_row(table) -> None:
    row = table.rows[0]
    _repeat_header(row)
    for cell, name in zip(row.cells, COLUMNS, strict=True):
        _set_width(cell, WIDTHS[name])
        _shade(cell, HEADER_FILL)
        p = cell.paragraphs[0]
        run = p.add_run(name)
        run.bold = True
        _tight(p, before=1, after=1)
        p.paragraph_format.keep_with_next = True  # the section heading, header row and first row stay together


def _parameter_cell(cell, p: dict, key: bool) -> None:
    symbol = clean(p.get("symbol"))
    code = (clean(p.get("code")) or "not in code").split(" (")[0]
    par = cell.paragraphs[0]
    run = par.add_run(("★ " if key else "") + (symbol or p["id"]))
    run.bold = True
    run.font.size = Pt(9)
    _tight(par)
    q = cell.add_paragraph()
    _runs(q, p["id"] + (" [derived]" if p.get("derived") else ""), color=MUTED)
    _tight(q)
    r = cell.add_paragraph()
    run = r.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(6.5)
    run.font.color.rgb = MUTED
    _tight(r)
    grp = p.get("group")
    if grp:
        g = cell.add_paragraph()
        _runs(g, f"group: {grp}", size=7, italic=True, color=MUTED)
        _tight(g)


def _value_cell(cell, value: str, treatment: str | None = None) -> None:
    par = cell.paragraphs[0]
    _runs(par, plain(value))
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _tight(par)
    if treatment:
        t = cell.add_paragraph()
        run = t.add_run(treatment)
        run.font.size = Pt(7)
        run.italic = True
        ink = TREATMENT_INK.get(treatment.split()[0])
        if ink:
            run.font.color.rgb = RGBColor.from_string(ink)
        _tight(t)


def _status_cell(cell, status: str, todo: str) -> None:
    par = cell.paragraphs[0]
    run = par.add_run(status)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(STATUS_INK.get(status, "5F6670"))
    _tight(par)
    if todo:
        t = cell.add_paragraph()
        _runs(t, "to-do: " + plain(todo), size=7, italic=True, color=MUTED)
        _tight(t)
    _shade(cell, STATUS_FILL.get(status, "FFFFFF"))


def render_docx(spec: dict):
    doc = Document()
    _page_setup(doc)
    title = doc.add_paragraph()
    run = title.add_run("Calibration parameters: US values and EU treatment")
    run.bold = True
    run.font.size = Pt(14)
    _tight(title, after=3)
    meta = spec.get("meta", {})
    intro = doc.add_paragraph()
    _runs(intro, f"{meta.get('paper', '')}. {meta.get('paper_table', '')}. ", color=MUTED)
    _runs(intro, "Generated from calibration_parameters.yaml by render_table.py; edit the YAML, not this file.", color=MUTED)
    _tight(intro, after=3)
    note = clean(meta.get("table_note"))
    if note:
        n = doc.add_paragraph()
        _runs(n, plain(note), color=MUTED)
        _tight(n, after=3)
    legend = doc.add_paragraph()
    _runs(legend, "★ key difference: where our EU reading departs from Antropic by judgement. "
          "Status: decided, discuss (value set, important and uncertain), open. "
          "Treatment: retain, recalculate (same method, EU data), recalibrate (new value by judgement), derived.", size=7.5, color=MUTED)
    _tight(legend, after=6)

    sections = by_section(spec)
    for sec, sec_title in spec["sections"].items():
        plist = sections.get(sec, [])
        if not plist:
            continue
        h = doc.add_paragraph()
        run = h.add_run(f"{sec}. {sec_title}")
        run.bold = True
        run.font.size = Pt(11)
        _tight(h, before=8, after=3)
        h.paragraph_format.keep_with_next = True

        table = doc.add_table(rows=1, cols=len(COLUMNS))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        _header_row(table)
        for p in plist:
            cells = rows_for(p)
            status = cells[-1]
            key = bool(p["eu"].get("key_difference"))
            row = table.add_row()
            _no_split(row)
            for cell, name in zip(row.cells, COLUMNS, strict=True):
                _set_width(cell, WIDTHS[name])
                if status == "discuss" and name != "Status":
                    _shade(cell, DISCUSS_ROW_FILL)
            c = row.cells
            _parameter_cell(c[0], p, key)
            _fill_cell(c[1], cells[2])
            _value_cell(c[2], cells[3])
            _fill_cell(c[3], cells[4])
            _value_cell(c[4], cells[6], cells[5])
            _fill_cell(c[5], cells[7])
            _fill_cell(c[6], cells[8])
            _status_cell(c[7], status, clean(p["eu"].get("todo")))
    return doc


def write_docx(spec: dict, path: Path) -> None:
    render_docx(spec).save(str(path))


def main() -> None:
    spec = yaml.safe_load(SRC.read_text(encoding="utf-8"))
    out = SRC.with_suffix(".docx")
    write_docx(spec, out)
    print(f"wrote {out.name}, {len(spec['parameters'])} parameters")


if __name__ == "__main__":
    main()
